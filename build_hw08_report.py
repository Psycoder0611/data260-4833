# ruff: noqa: E501
"""Build the HW8 report (Word + PDF) from the raw files of my runs.

  python build_hw08_report.py                  Git copy: Report_editable.docx + Report.pdf
  python build_hw08_report.py --commit <hash>  also the Canvas copy:
                                               Kulshrestha_HW8_editable.docx + Kulshrestha_HW8.pdf
                                               (adds the tag hw8 and the commit hash)
  python build_hw08_report.py --list-screenshots

The text is written by me. Numbers in the tables are read from reports/hw08/raw/
(JMeter CSV exports, part3_*.json). Anything not run yet shows "[run pending]".
Screenshots go in reports/hw08/screenshots/ as fig01_*.png, fig02_*.png ...
(see --list-screenshots). A missing screenshot becomes an empty labelled slot, so
I can also paste the picture into the Word file by hand and export the PDF myself.
"""

import argparse
import ast
import json
import shutil
import subprocess
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from hw8 import config, jtl

ROOT = Path(__file__).resolve().parent
OUT = config.REPORT_DIR
RAW = config.RAW_DIR
SHOTS = OUT / "screenshots"
REPO_URL = "https://github.com/Psycoder0611/data260-4833"
PENDING = "[run pending]"
FONT = "Calibri"

# (key, caption, how to capture it)
FIGURES = [
    (
        "docker_ps",
        "docker ps: Kafka, worker and API containers running",
        "Terminal: `docker ps` after `make up-hw08` (name visible in the prompt).",
    ),
    (
        "worker_ready",
        "Worker log: Consumer is ready and listening...",
        "Terminal: `docker compose -f docker-compose.hw8.yml logs worker`.",
    ),
    (
        "api_started",
        "FastAPI log: started at http://127.0.0.1:8333",
        "Terminal: `docker compose -f docker-compose.hw8.yml logs api`.",
    ),
    (
        "demo_steps",
        "homework_demo.py: the ten logged steps, success path",
        "Terminal: `python homework_demo.py` (top part, steps 1-10).",
    ),
    (
        "demo_errors",
        "homework_demo.py: unknown operation, bad data and unknown id",
        "Same run: the EXTRA section with the error cases.",
    ),
    (
        "api_postman",
        "POST /api/hw8/entity through Kafka returns the response",
        "curl or Postman screenshot of one POST to http://127.0.0.1:8333/api/hw8/entity.",
    ),
    (
        "agents_ready",
        "Planner, Writer and Reviewer running as three processes",
        "Three terminal windows side by side, each showing Consumer is ready and listening...",
    ),
    (
        "journey",
        "One question travelling inbox -> tasks -> drafts -> final",
        "Terminal: `python send_question.py --wait --journey`.",
    ),
    (
        "final_json",
        "Final JSON with status and answer",
        "Same command: the final JSON block at the end.",
    ),
    (
        "dup_before",
        "Duplicate delivery without idempotency: two drafts",
        "Terminal: `python run_hw08_part3.py duplicate` (BEFORE part).",
    ),
    (
        "dup_after",
        "Duplicate delivery with idempotency: one draft, replay skipped",
        "Same run (AFTER part, DUPLICATE line).",
    ),
    (
        "death_before",
        "Consumer killed, commit on receive: message lost",
        "Terminal: `python run_hw08_part3.py death` (BEFORE part).",
    ),
    (
        "death_after",
        "Consumer killed, commit after success: message redelivered",
        "Same run (AFTER part, RECEIVED again then COMMITTED).",
    ),
    (
        "poison_before",
        "Poison message without a limit: endless retries, queue blocked",
        "Terminal: `python run_hw08_part3.py poison` (BEFORE part).",
    ),
    (
        "poison_after",
        "Poison message with 3 attempts and the dead-letter topic",
        "Same run (AFTER part, DEAD LETTER line and next message processed).",
    ),
    (
        "ordering_unkeyed",
        "Ordering without a key: out-of-order result table",
        "Terminal: `python run_hw08_part3.py ordering` (unkeyed table).",
    ),
    (
        "ordering_keyed",
        "Ordering with the entity as key: in order",
        "Same run (keyed table).",
    ),
    (
        "jmeter_kafka_plan",
        "JMeter plan for the Kafka endpoint (threads, ramp-up, loop)",
        "JMeter GUI with kafka_endpoint.jmx open; show the Thread Group settings.",
    ),
    (
        "jmeter_kafka_results",
        "JMeter Kafka endpoint: Aggregate Report (100 samples, 0% errors)",
        "JMeter GUI after the run: Summary Report / Aggregate Report listeners.",
    ),
    (
        "jmeter_direct_results",
        "JMeter direct endpoint: Aggregate Report (100 samples, 0% errors)",
        "Same for direct_endpoint.jmx.",
    ),
    (
        "jmeter_graph",
        "JMeter Response Time Graph for the Kafka endpoint (1 s intervals)",
        "Response Time Graph listener and one request in View Results Tree.",
    ),
    (
        "verify",
        "make verify-hw08: all checks pass",
        "Terminal: `make verify-hw08` (the last lines with OVERALL: PASS).",
    ),
]


# code screenshots: key -> (file prefix, caption, how to capture it)
CODE = {
    "code_validate": ("code02", "Code: validate_create (hw8/entity_service.py)", ""),
    "code_process": ("code03", "Code: process_request (hw8/entity_service.py)", ""),
    "code_worker": ("code01", "Code: worker handle (hw8/entity_worker.py)", ""),
    "code_rpc": ("code04", "Code: Kafka client request/wait (hw8/rpc_client.py)", ""),
    "code_api_kafka": ("code05", "Code: POST /api/hw8/entity over Kafka (hw8/api.py)", ""),
    "code_api_direct": ("code06", "Code: /api/hw8/entity-direct baseline (hw8/api.py)", ""),
    "code_demo": (
        "code14",
        "Code: homework_demo.py, demo() steps 1-10",
        "Screenshot of homework_demo.py lines ~106-160 (create correlation_id, CREATE request, "
        "extract entity_id, GET request, correlation_id check). Save as "
        "reports/hw08/screenshots/code14_homework_demo.png",
    ),
    "code_planner": ("code11", "Code: Planner (planner.py)", ""),
    "code_writer": ("code12", "Code: Writer (writer.py)", ""),
    "code_reviewer": ("code13", "Code: Reviewer (reviewer.py)", ""),
    "code_idem": ("code09", "Code: idempotency store (hw8/idempotency.py)", ""),
    "code_commit": ("code07", "Code: process_record, commit and duplicate check (hw8/agent_base.py)", ""),
    "code_retry": ("code08", "Code: bounded retries in _attempts (hw8/agent_base.py)", ""),
    "code_order": ("code10", "Code: partition key and message generator (hw8/ordering.py)", ""),
}

# figure order = order in the report (code and output together, section by section)
ORDER = [
    "code_validate", "code_process",
    "code_worker", "docker_ps", "worker_ready",
    "code_api_kafka", "code_api_direct", "api_started", "api_postman",
    "code_demo", "code_rpc", "demo_steps", "demo_errors",
    "code_planner", "code_writer", "code_reviewer", "agents_ready", "journey", "final_json",
    "code_idem", "dup_before", "dup_after",
    "code_commit", "death_before", "death_after",
    "code_retry", "poison_before", "poison_after",
    "code_order", "ordering_unkeyed", "ordering_keyed",
    "jmeter_kafka_plan", "jmeter_kafka_results", "jmeter_direct_results", "jmeter_graph",
    "verify",
]  # fmt: skip
PLACED: List[str] = []


# --------------------------------------------------------------------- helpers
def fig_no(key: str) -> int:
    """Figure number of a screenshot key (position in the report)."""
    return ORDER.index(key) + 1


def ref(*keys: str) -> str:
    """'Figure 3' or 'Figures 3 and 4' for the given keys."""
    nums = [fig_no(k) for k in keys]
    if len(nums) == 1:
        return f"Figure {nums[0]}"
    return "Figures " + ", ".join(str(n) for n in nums[:-1]) + f" and {nums[-1]}"


def load(name: str) -> Optional[Dict[str, Any]]:
    """Load a raw JSON file, or None."""
    path = RAW / name
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def source_of(path: str, func: str, cls: Optional[str] = None) -> str:
    """Source text of a function (or method) from the repo, found with ast."""
    text = (ROOT / path).read_text(encoding="utf-8")
    tree = ast.parse(text)
    scope: List[ast.AST] = [tree]
    if cls:
        scope = [
            n for n in ast.walk(tree) if isinstance(n, ast.ClassDef) and n.name == cls
        ]
    for container in scope:
        for node in ast.walk(container):
            if (
                isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
                and node.name == func
            ):
                lines = text.splitlines()[node.lineno - 1 : node.end_lineno]
                indent = len(lines[0]) - len(lines[0].lstrip())
                return "\n".join(line[indent:] for line in lines)
    return f"# {func} not found in {path}"


# ------------------------------------------------------------------- docx tools
class Doc:
    """Small wrapper that keeps the formatting rules in one place."""

    def __init__(self) -> None:
        """Create the document with Calibri 11 pt body text."""
        self.d = Document()
        section = self.d.sections[0]
        section.left_margin = section.right_margin = Inches(0.9)
        section.top_margin = section.bottom_margin = Inches(0.8)
        style = self.d.styles["Normal"]
        style.font.name = FONT
        style.font.size = Pt(11)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_after = Pt(6)

    def _run(
        self,
        p,
        text: str,
        size: float,
        bold: bool = False,
        mono: bool = False,
        color: Optional[str] = None,
        italic: bool = False,
    ):
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.name = "Consolas" if mono else FONT
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        return run

    def title(self, text: str) -> None:
        """Big centred title."""
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._run(p, text, 24)

    def centered(self, text: str, size: float = 13) -> None:
        """Centred line of text."""
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._run(p, text, size, color="555555")

    def h1(self, text: str) -> None:
        """Main heading: 14 pt bold."""
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.keep_with_next = True
        self._run(p, text, 14, bold=True)

    def h2(self, text: str) -> None:
        """Subheading: 12 pt bold."""
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = True
        self._run(p, text, 12, bold=True)

    def p(self, text: str) -> None:
        """Body paragraph: 11 pt."""
        p = self.d.add_paragraph()
        self._run(p, text, 11)

    def label(self, label: str, value: str) -> None:
        """Bold label followed by normal text (11 pt)."""
        p = self.d.add_paragraph()
        self._run(p, label, 11, bold=True)
        self._run(p, value, 11)

    def bullets(self, items: List[str]) -> None:
        """Bulleted list (11 pt)."""
        for item in items:
            p = self.d.add_paragraph(style="List Bullet")
            self._run(p, item, 11)

    def code(self, text: str, caption: str) -> None:
        """Monospace block with a small caption above it."""
        c = self.d.add_paragraph()
        c.paragraph_format.keep_with_next = True
        self._run(c, caption, 9.5, italic=True, color="555555")
        p = self.d.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:fill"), "F2F2F2")
        p._p.get_or_add_pPr().append(shade)
        self._run(p, text, 8.5, mono=True)

    def table(
        self,
        rows: List[List[Any]],
        widths: Optional[List[float]] = None,
        header: bool = True,
    ) -> None:
        """Table with 9.5 pt text and a bold header row."""
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                self._run(cell.paragraphs[0], str(value), 9.5, bold=header and i == 0)
                cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                # keep the whole table on one page: every row but the last
                # stays with the next row
                if i < len(rows) - 1:
                    cell.paragraphs[0].paragraph_format.keep_with_next = True
                if widths:
                    cell.width = Inches(widths[j])
            row_props = t.rows[i]._tr.get_or_add_trPr()
            row_props.append(OxmlElement("w:cantSplit"))
        self.d.add_paragraph()

    def figure(self, key: str) -> None:
        """A screenshot (or an empty labelled slot) with its caption."""
        number = fig_no(key)
        PLACED.append(key)
        if key in CODE:
            prefix, caption, how = CODE[key]
            matches = sorted(SHOTS.glob(f"{prefix}_*.png"))
        else:
            idx = [f[0] for f in FIGURES].index(key) + 1
            caption, how = FIGURES[idx - 1][1], FIGURES[idx - 1][2]
            matches = sorted(SHOTS.glob(f"fig{idx:02d}_*")) + sorted(
                SHOTS.glob(f"fig{idx:02d}.*")
            )
        if matches:
            self.d.add_picture(str(matches[0]), width=Inches(6.2))
            self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.d.paragraphs[-1].paragraph_format.keep_with_next = True
        else:
            t = self.d.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            cell = t.cell(0, 0)
            cell.text = ""
            self._run(
                cell.paragraphs[0],
                f"[Paste screenshot for Figure {number} here]\n{how}",
                10,
                color="888888",
                italic=True,
            )
            for _ in range(5):
                cell.add_paragraph()
        cap = self.d.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._run(cap, f"Figure {number} - {caption}", 9.5, color="333333")

    def figs(self, *keys: str) -> None:
        """Several figures one below the other."""
        for k in keys:
            self.figure(k)

    def page_break(self) -> None:
        """Start a new page."""
        self.d.add_page_break()

    def save(self, path: Path) -> None:
        """Write the .docx."""
        self.d.save(path)


# ------------------------------------------------------------- data for tables
def jmeter_rows() -> List[List[Any]]:
    """Rows for the JMeter comparison table (from the CSV exports)."""
    rows: List[List[Any]] = [
        [
            "Endpoint",
            "Samples",
            "Avg (ms)",
            "Min (ms)",
            "Max (ms)",
            "p95 (ms)",
            "Error %",
            "Req/s",
        ]
    ]
    for label, name in (
        ("Kafka request/response", "jmeter_kafka.csv"),
        ("Direct (no Kafka)", "jmeter_direct.csv"),
    ):
        path = RAW / name
        if not path.exists():
            rows.append([label] + [PENDING] + [""] * 6)
            continue
        s = jtl.summarize(path)
        tp = s["throughput_per_sec"] if s["throughput_per_sec"] is not None else "n/a"
        rows.append(
            [
                label,
                s["samples"],
                s["average_ms"],
                s["min_ms"],
                s["max_ms"],
                s["p95_ms"],
                f"{s['error_pct']}%",
                tp,
            ]
        )
    return rows


def jmeter_comment() -> str:
    """One sentence of reading, written from the numbers (or a prompt to write it)."""
    k, d = RAW / "jmeter_kafka.csv", RAW / "jmeter_direct.csv"
    if not (k.exists() and d.exists()):
        return "[Write 3-4 sentences here after the JMeter runs: which endpoint was slower, by how much, and why.]"
    ks, ds = jtl.summarize(k), jtl.summarize(d)
    gap = ks["average_ms"] - ds["average_ms"]
    return (
        f"On average the Kafka endpoint took {ks['average_ms']} ms per request and the direct endpoint {ds['average_ms']} ms, "
        f"so the trip through Kafka added about {gap:.0f} ms per request. Error rates were {ks['error_pct']}% (Kafka) and {ds['error_pct']}% (direct)."
    )


def failure_rows() -> List[List[Any]]:
    """The failure table with before/after evidence."""
    rows: List[List[Any]] = [["Failure", "Before the fix", "After the fix", "Evidence"]]
    dup, death, poison = (
        load("part3_duplicate.json"),
        load("part3_death.json"),
        load("part3_poison.json"),
    )
    if dup:
        rows.append(
            [
                "Duplicate delivery",
                f"Same message sent twice -> {dup['before']['drafts_produced']} drafts (work done twice)",
                f"{dup['after']['drafts_produced']} draft; replay logged as DUPLICATE and skipped ({dup['after']['duplicates_skipped']} skipped)",
                f"{ref('dup_before', 'dup_after')}",
            ]
        )
    else:
        rows.append(
            ["Duplicate delivery", PENDING, PENDING, ref("dup_before", "dup_after")]
        )
    if death:
        b, a = death["before"], death["after"]
        rows.append(
            [
                "Consumer dies mid-work",
                f"Offset committed on receive; after restart redelivered = {b['victim_redelivered_after_restart']}, draft produced = {b['victim_draft_produced']} (message lost)",
                f"Offset committed after success; redelivered = {a['victim_redelivered_after_restart']}, draft produced = {a['victim_draft_produced']}",
                f"{ref('death_before', 'death_after')}",
            ]
        )
    else:
        rows.append(
            [
                "Consumer dies mid-work",
                PENDING,
                PENDING,
                ref("death_before", "death_after"),
            ]
        )
    if poison:
        b, a = poison["before"], poison["after"]
        rows.append(
            [
                "Poison message",
                f"Retried without limit ({b['failed_attempts']} attempts seen in the window); next message processed = {b['good_message_processed']}",
                f"{a['failed_attempts']} attempts, then 1 message in s4833_dead_letter; next message processed = {a['good_message_processed']}",
                f"{ref('poison_before', 'poison_after')}",
            ]
        )
    else:
        rows.append(
            ["Poison message", PENDING, PENDING, ref("poison_before", "poison_after")]
        )
    order = load("part3_ordering.json")
    if order:
        u, k = order["unkeyed"]["analysis"], order["keyed"]["analysis"]
        rows.append(
            [
                "Two partitions / two consumers",
                f"No key: related messages split over both partitions; {u['total_inversions']} order inversions, {u['entities_out_of_order']} entities out of order, {u['entities_in_more_than_one_partition']} entities spread over both partitions",
                f"Entity id as key: {k['total_inversions']} inversions, {k['entities_out_of_order']} entities out of order, all messages of an entity stay in one partition",
                f"{ref('ordering_unkeyed', 'ordering_keyed')}",
            ]
        )
    else:
        rows.append(
            ["Two partitions / two consumers", PENDING, PENDING, ref("ordering_unkeyed", "ordering_keyed")]
        )
    return rows


def ordering_rows() -> List[List[Any]]:
    """Ordering results with and without a key."""
    rows: List[List[Any]] = [
        [
            "Run",
            "Messages processed",
            "Order inversions",
            "Entities out of order",
            "Entities split over both partitions",
        ]
    ]
    data = load("part3_ordering.json")
    for mode, label in (
        ("unkeyed", "No key (random partition)"),
        ("keyed", "Key = entity"),
    ):
        if not data:
            rows.append([label, PENDING, "", "", ""])
            continue
        a = data[mode]["analysis"]
        rows.append(
            [
                label,
                a["messages_processed"],
                a["total_inversions"],
                a["entities_out_of_order"],
                a["entities_in_more_than_one_partition"],
            ]
        )
    return rows


def ordering_comment() -> str:
    """Describe the unkeyed result as it really came out in my run."""
    data = load("part3_ordering.json")
    if not data:
        return "[Describe the no-key result after the run.]"
    a = data["unkeyed"]["analysis"]
    spread = a["entities_in_more_than_one_partition"]
    if a["total_inversions"] > 0:
        return (
            f"In my run {a['entities_out_of_order']} of 5 entities were processed out of order "
            f"({a['total_inversions']} inversions in total) and {spread} entities had messages "
            "in both partitions."
        )
    return (
        f"In my run all {spread} entities were spread over both partitions, but the two Writers "
        "happened to finish each entity's messages in the right order, so there were 0 inversions. "
        "That was luck: without a key Kafka promises no order across partitions, and the key run "
        "below is the one that makes it a guarantee."
    )


# ----------------------------------------------------------------------- report
def build(canvas: bool, commit: str) -> Doc:
    """Write the whole report text."""
    PLACED.clear()
    d = Doc()
    d.title("DATA 260: Agentic AI")
    d.centered(
        "Instructor: Prof. Dr. Simon Shim   |   Student Name: Trapti Kulshrestha   |   STUDENT ID: 019124833   |   Homework: 8"
    )
    d.label("GitHub Repository: ", REPO_URL)

    d.h1("Section 0: Personal Configuration")
    d.table([
        ["Setting", "Value"],
        ["SID4", "4833"],
        ["PORT_BASE", "8333 (8000 + 4833 mod 900)"],
        ["PREFIX", "s4833"],
        ["SEED", "4833"],
        ["VERIFY_SEED", "264833"],
        ["DOMAIN_ID", "1 (4833 mod 8)"],
        ["Assigned Domain", "Clinical Trial Listings"],
    ], [2.2, 4.4])  # fmt: skip
    d.label("Hardware: ", "Apple MacBook Pro")
    d.label("Local Model: ", "qwen3:8b (Ollama), seed 4833. No hosted API is used.")
    d.label(
        "Message broker: ",
        "Apache Kafka 3.9.2 (KRaft mode) in Docker; Python client aiokafka 0.14.0",
    )
    if canvas:
        d.label("Tagged commit (tag hw8): ", commit)

    d.h1("Project Setup")
    d.p(
        "For Homework 8 I kept working in my existing DATA 260 repository, on a hw8-work branch made from the hw7 tag. All new code is in the hw8/ package and a few top-level scripts, so none of the earlier homework is changed. Kafka runs from docker-compose.hw8.yml (one broker, KRaft, so no ZooKeeper). A small init container creates the seven topics with the partition counts I need and automatic topic creation is switched off, so a typo in a topic name fails loudly instead of creating a new topic. The worker and the FastAPI wrapper are two more containers built from the same image. The Planner, Writer and Reviewer for Part 2 run on my Mac, outside Docker, because they call Ollama on the host."
    )
    d.p(
        "Topics (all prefixed s4833): entity_requests, entity_responses, inbox, drafts, final and dead_letter have one partition each; tasks has two partitions, because Part 3 needs two Writers to share it."
    )

    d.h1("Part 1 - Kafka Request/Response")
    d.h2("1.1 Message format and validation")
    d.p(
        "A client sends a request to s4833_entity_requests and waits for the answer on s4833_entity_responses. A request has four fields: correlation_id (a UUID made by the client), reply_to (the topic the answer must go to), operation and data. The worker supports two operations for my domain, CREATE_CLINICAL_TRIAL and GET_CLINICAL_TRIAL. A create needs trialTitle, sponsor, trial_code, participant_capacity (an integer of at least 1), submitterEmail and category; description and organization_id are optional."
    )
    d.p(
        "The worker never crashes on bad input. If the envelope is wrong (no correlation_id, no reply_to, data that is not an object, a message that is not JSON) it logs it and, when it can still find a reply topic, answers with success false and an error. If the operation is unknown, the data is invalid or the trial id does not exist, the answer is still a normal response with success false, a short error text and the same correlation_id, so the client can match the error to its request. The two screenshots below show the validation code and process_request, which turns every request into a response."
    )
    d.figs("code_validate", "code_process")
    d.h2("1.2 The worker")
    d.p(
        "The worker joins the consumer group s4833_entity_worker and logs 'Consumer is ready and listening...' once partitions are assigned. For every request it logs REQUEST RECEIVED, PROCESSING, PROCESSED and RESPONSE SENT with the correlation_id, so one request can be followed through the log with a search. It commits the Kafka offset only after the response has been sent. If the same correlation_id arrives a second time, the worker sends back the stored first response instead of creating a second trial. "
        + f"{ref('docker_ps', 'worker_ready')} show the containers and the ready message. "
        + ref("worker_ready")
        + " was taken after restarting the worker container: the last lines of the log are a few earlier requests, then the ready line with the request and response topics."
    )
    d.figs("code_worker", "docker_ps", "worker_ready")
    d.h2("1.3 FastAPI wrapper")
    d.p(
        "The FastAPI app turns an HTTP POST into a Kafka request and waits for the matching response. It listens on port 8333 and logs 'FastAPI server started at http://127.0.0.1:8333' on start-up ("
        + ref("api_started")
        + "). If no answer arrives in time it returns HTTP 504. A second route, /api/hw8/entity-direct, runs the same business logic without Kafka; I added it only as the baseline for the JMeter comparison."
    )
    d.figs("code_api_kafka", "code_api_direct", "api_started", "api_postman")
    d.h2("1.4 homework_demo.py")
    d.p(
        "homework_demo.py shows the ten steps the assignment lists: client starts, builds the request, generates the UUID, publishes, the worker receives it, validates, processes, sends the response, the client matches the correlation_id and prints the result. The worker's own log lines are merged into the same output so the whole round trip reads in order. "
        + ref("demo_steps")
        + " is the success path. After the ten steps it also sends an unknown operation, invalid data and an unknown id; each returns success false with the original correlation_id ("
        + ref("demo_errors")
        + "). "
        + ref("api_postman")
        + " shows the same round trip over HTTP. The two code screenshots above them are the demo steps and the Kafka client (rpc_client.py) that sends the request and waits for the matching correlation_id."
    )
    d.figs("code_demo", "code_rpc", "demo_steps", "demo_errors")

    d.h1("Part 2 - Planner, Writer and Reviewer over Kafka")
    d.p(
        "The three agents are separate Python processes (planner.py, writer.py, reviewer.py). Each one reads from one topic and writes to the next: a question goes into s4833_inbox, the Planner writes a plan to s4833_tasks, the Writer writes a draft to s4833_drafts, and the Reviewer scores it and writes the result to s4833_final. The correlation_id is created once in send_question.py and copied unchanged by every agent, so one grep follows the whole journey. The Planner, Writer and Reviewer call qwen3:8b through the HW1 model client with seed 4833; the calls run in a worker thread so the Kafka heartbeat keeps going while the model is thinking. "
        + ref("agents_ready")
        + " shows the three processes ready."
    )
    d.p(
        "send_question.py --wait --journey sends one question and then reads all four topics to print every hop with its partition and offset ("
        + ref("journey")
        + "). The final message holds the status (approved or rejected, from a review score of at least 6 out of 10) and the answer ("
        + ref("final_json")
        + ")."
    )
    d.figs("code_planner", "code_writer", "code_reviewer", "agents_ready", "journey", "final_json")

    d.h1("Part 3 - Failure Handling")
    d.p(
        "For each failure I first ran the agent with the protection switched off to see the problem, then switched the protection on and ran the same test again. The switches are command-line options on the same Writer code (--no-idempotency, --commit-mode on_receive, --max-attempts 0 --no-dead-letter), so the before and after runs differ only by the fix. Everything is started by python run_hw08_part3.py, and the Writer in these runs uses the fixed-text stand-in instead of qwen3:8b (see Self-check)."
    )
    d.h2("3.1 Duplicate delivery and idempotency")
    d.p(
        "Kafka gives at-least-once delivery, so the same message can arrive twice. I sent one task message and then sent the same message again with the same correlation_id. Without a check the Writer did the work twice and published two drafts ("
        + ref("dup_before")
        + "). The fix is a small SQLite table of (stage, correlation_id) pairs that are already done. The Writer checks it before working and records the id after it has published the draft; a replay is logged as DUPLICATE, skipped, and its offset is committed so it does not come back ("
        + ref("dup_after")
        + ")."
    )
    d.figs("code_idem", "dup_before", "dup_after")
    d.h2("3.2 Consumer killed in the middle of work")
    d.p(
        "I gave the victim message a long processing time, killed the Writer with SIGKILL while it was working, and started a new Writer in the same group. If the offset is committed when the message is received, the new Writer starts after it and the message is lost: no draft is ever produced ("
        + ref("death_before")
        + "). With commit-after-success the offset is only committed once the draft has been published, so the dead Writer never committed it, the new Writer receives it again, finishes it and then commits ("
        + ref("death_after")
        + "). The cost is that a message can be processed twice if the Writer dies after publishing but before committing, which is why Part 3.1 and this part belong together."
    )
    d.figs("code_commit", "death_before", "death_after")
    d.h2("3.3 Poison message and the dead-letter topic")
    d.p(
        "A poison message is one that can never succeed. Mine has a plan field that is a number instead of a list, so the Writer raises an error every time. With unlimited retries the Writer keeps failing on the same offset and the good message behind it is never reached ("
        + ref("poison_before")
        + "). With the fix the Writer tries three times with a growing wait, then publishes the original message, the error, the number of attempts and its source topic, partition and offset to s4833_dead_letter, commits the offset and moves on. The good message after it is processed normally ("
        + ref("poison_after")
        + ")."
    )
    d.figs("code_retry", "poison_before", "poison_after")
    d.h2("3.4 Ordering")
    d.p(
        "I generated 20 messages with VERIFY_SEED 264833: five clinical-trial entities with four messages each, where the order of the four matters. s4833_tasks has two partitions and two Writers run in the same consumer group, so each Writer owns one partition. First I sent the messages with no key; Kafka then spreads them over both partitions, the two Writers work at different speeds, so the order inside one entity is not guaranteed ("
        + ref("ordering_unkeyed")
        + "). "
        + ordering_comment()
        + " Then I sent the same kind of messages with the entity id as the key. The key decides the partition, so all messages of one entity go to the same partition, and one partition is read by one Writer at a time, in order ("
        + ref("ordering_keyed")
        + ")."
    )
    d.figs("code_order", "ordering_unkeyed", "ordering_keyed")
    d.table(ordering_rows(), [1.9, 1.1, 1.1, 1.2, 1.5])
    d.p(
        "The key only keeps order inside one entity. Different entities can still finish in any order, which is fine for this homework because nothing depends on it."
    )

    d.h1("Failure Table")
    d.table(failure_rows(), [1.2, 2.2, 2.2, 1.0])

    d.h1("Comparison with the HW2 In-Process Graph")
    d.p(
        "In HW2 my Planner, Reviewer and Supervisor were nodes of one LangGraph in a single Python process. They shared an AgentState dictionary, called each other as function calls and a turn ceiling in the Supervisor stopped endless loops. In HW8 the agents are separate processes that only share messages."
    )
    d.h2("Two benefits")
    d.bullets(
        [
            "Independence and scaling. Each agent can crash, be restarted or be run twice without the others noticing. In Part 3.4 I ran two Writers in one group and Kafka split the partitions between them. In HW2 one exception in a node stopped the whole graph run.",
            "A durable record. Every hop is stored in a topic with a partition and an offset, so I can replay a message or look at exactly what the Planner sent. In HW2 the state existed only in memory while the run was alive.",
        ]
    )
    d.h2("Two new responsibilities")
    d.bullets(
        [
            "Delivery has to be handled by me. A function call either happens or raises; a message can be delivered twice or lost. I had to add idempotency, commit-after-success and a bounded retry with a dead-letter topic (Parts 3.1 to 3.3) just to get back the safety the graph gave for free.",
            "Ordering and operations. In the graph the order was the order of the edges. With partitions I must choose a key to get order, and I must also run and watch Kafka itself (topics, partitions, consumer groups, the broker container).",
        ]
    )
    d.p(
        "There is also a cost in speed, shown in the JMeter section below. One thing I did not rebuild: the HW2 Supervisor could send a rejected draft back to the Planner. In HW8 the Reviewer only marks the draft approved or rejected; there is no loop back."
    )

    d.h1("JMeter Load Test")
    d.p(
        "I compared the two FastAPI routes with Apache JMeter: /api/hw8/entity (through Kafka and the worker) and /api/hw8/entity-direct (no Kafka). Both plans use 100 threads, a 10 second ramp-up and 1 loop, an HTTP POST with a JSON body (the trial_code contains the thread number so every request creates a different trial) and a Content-Type: application/json header. A response assertion checks for success true. Each plan has the Results Tree, Summary Report, Aggregate Report and Response Time Graph listeners. The numbers below are exported from the listeners (Save Table Data) and summarised by make_hw08_metrics.py. The plan settings are in "
        + ref("jmeter_kafka_plan")
        + ", the Kafka results in "
        + ref("jmeter_kafka_results")
        + ", the direct results in "
        + ref("jmeter_direct_results")
        + " and the graph in "
        + ref("jmeter_graph")
        + "."
    )
    d.figs("jmeter_kafka_plan", "jmeter_kafka_results", "jmeter_direct_results", "jmeter_graph")
    d.table(jmeter_rows(), [1.7, 0.7, 0.7, 0.7, 0.7, 0.8, 0.7, 0.7])
    d.p(jmeter_comment())

    d.h1("Self-check")
    d.p(
        "make verify-hw08 runs verify_hw08.py. It checks the topics, a create-then-get round trip with the correlation_id kept, the error paths, the FastAPI routes, one question through the pipeline, the three failure fixes, the ordering key, the offline tests and that the JMeter plans are present, and writes reports/hw08/verification.json ("
        + ref("verify")
        + "). The Part 3 failure runs (and the checks in this script) use a fixed-text stand-in instead of qwen3:8b for the Writer, because what they test is delivery, offsets and ordering, not the quality of the text, and a stand-in makes the runs short and repeatable. Part 2 uses qwen3:8b."
    )
    d.figs("verify")

    d.h1("AI Use")
    ai = (
        (OUT / "AI_USE.md").read_text(encoding="utf-8")
        if (OUT / "AI_USE.md").exists()
        else ""
    )
    for block in ai.split("\n## ")[1:]:
        heading, _, body = block.partition("\n")
        d.h2(heading.strip())
        for para in body.split("\n\n"):
            para = para.strip()
            if para and not para.startswith("<!--"):
                d.p(para)

    assert PLACED == ORDER, f"figure order mismatch: {PLACED}"
    return d


def to_pdf(docx_path: Path) -> Optional[Path]:
    """Convert with LibreOffice if it is installed."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found: open the .docx in Word and export to PDF.")
        return None
    subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx_path.parent),
            str(docx_path),
        ],
        check=True,
        capture_output=True,
    )
    return docx_path.with_suffix(".pdf")


def main() -> None:
    """Build the Git copy, and the Canvas copy when --commit is given."""
    parser = argparse.ArgumentParser(description="Build the HW8 report")
    parser.add_argument(
        "--commit", default="", help="hw8 tag commit hash (Canvas copy)"
    )
    parser.add_argument("--list-screenshots", action="store_true")
    parser.add_argument("--no-pdf", action="store_true")
    args = parser.parse_args()
    if args.list_screenshots:
        for i, (key, caption, how) in enumerate(FIGURES, 1):
            have = bool(list(SHOTS.glob(f"fig{i:02d}*")))
            print(
                f"{'[x]' if have else '[ ]'} fig{i:02d}_{key}.png  {caption}\n        {how}"
            )
        return
    OUT.mkdir(parents=True, exist_ok=True)
    jobs = [(False, "", "Report_editable.docx", "Report.pdf")]
    if args.commit:
        jobs.append(
            (True, args.commit, "Kulshrestha_HW8_editable.docx", "Kulshrestha_HW8.pdf")
        )
    for canvas, commit, docx_name, pdf_name in jobs:
        path = OUT / docx_name
        build(canvas, commit).save(path)
        print(f"wrote {path}")
        if not args.no_pdf:
            made = to_pdf(path)
            if made and made.name != pdf_name:
                made.replace(OUT / pdf_name)
            if made:
                print(f"wrote {OUT / pdf_name}")


if __name__ == "__main__":
    main()
